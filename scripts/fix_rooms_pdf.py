import os
import sys
import json
import openpyxl
import re
import html
import base64
import subprocess

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output_pdf")
WEB_DIR = os.path.join(BASE_DIR, "web")

logo_path = os.path.join(WEB_DIR, 'images.png')
with open(logo_path, 'rb') as f:
    logo_base64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

map_path = os.path.join(DATA_DIR, 'clean_room_building_map.json')
mapping = {}
if os.path.exists(map_path):
    with open(map_path, 'r', encoding='utf-8') as f:
        mapping = json.load(f)
pairs = mapping.get('pairs', {})
codes = mapping.get('codes', {})

def get_building_name(kd, nm):
    kd_u = str(kd or '').strip().upper()
    nm_l = str(nm or '').strip().lower()
    if not kd_u and not nm_l:
        return ''
    key = f"{kd_u.lower()}|||{nm_l}"
    if key in pairs:
        return pairs[key]
    if kd_u.startswith('D.1.') or kd_u.startswith('D.2.') or kd_u.startswith('D.'):
        return 'Asrama Marsawa'
    if kd_u.startswith('F.1.') or kd_u.startswith('F.2.') or kd_u.startswith('F.'):
        return 'Asrama Bulian'
    if kd_u.startswith('E.1.') or kd_u.startswith('E.'):
        return 'Asrama Kulim'
    if kd_u.startswith('G.1.') or kd_u.startswith('G.2.') or kd_u.startswith('G.'):
        return 'Mess Jelutung 3'
    if kd_u.startswith('H.1.') or kd_u.startswith('H.'):
        return 'Mess Penatar'
    if (kd_u == '2' or kd_u == '3') and 'kamar' in nm_l:
        return 'Mess Kepala'
    if (kd_u == '4' or kd_u == '5') and 'kamar' in nm_l:
        return 'Mess Jelutung 2'
    kd_l = kd_u.lower()
    if kd_l not in ['1', '2', '3', '4', '5', '01', '02', '03', '04', '05'] and kd_l in codes:
        return codes[kd_l]
    return ''

def sanitize_filename(name, max_len=70):
    s = re.sub(r'[\<\>\:\"\/\\\\\|\?\*]', ' - ', str(name or ''))
    s = re.sub(r'\s+', ' ', s).strip('. ')
    if len(s) > max_len:
        s = s[:max_len].strip('. ')
    return s or 'Item'

def extract_year(tahun_val, tgl_val, satker_val):
    t_str = str(tahun_val or '').strip()
    m = re.search(r'\b(19\d{2}|20\d{2})\b', t_str)
    if m:
        return m.group(1)

    if isinstance(tgl_val, (int, float)) and 25000 < tgl_val < 60000:
        try:
            from datetime import datetime, timedelta
            base = datetime(1899, 12, 30)
            d = base + timedelta(days=tgl_val)
            return str(d.year)
        except Exception:
            pass

    d_str = str(tgl_val or '').strip()
    m2 = re.search(r'\b(19\d{2}|20\d{2})\b', d_str)
    if m2:
        return m2.group(1)

    s_str = str(satker_val or '').strip()
    m3 = re.search(r'\.(19\d{2}|20\d{2})\b', s_str)
    if m3:
        return m3.group(1)
    return ''

def render_label_html(item):
    satker = str(item.get('kode_satker') or '138041000693242000KD').strip()
    tahun = str(item.get('tahun') or '').strip()
    if not tahun and item.get('tgl_perolehan'):
        m = re.search(r'\b(19\d{2}|20\d{2})\b', str(item.get('tgl_perolehan')))
        if m:
            tahun = m.group(1)

    if tahun:
        if satker.endswith(f".{tahun}") or f".{tahun}" in satker:
            satker_display = satker
        else:
            satker_display = f"{satker}.{tahun}"
    else:
        satker_display = satker

    kode_barang = html.escape(str(item.get('kode_barang') or '-'))
    nup = html.escape(str(item.get('nup') or '1'))
    nama_barang = html.escape(str(item.get('nama_barang') or 'Barang BMN'))
    merk_tipe = html.escape(str(item.get('merk_tipe') or '-'))

    qr = str(item.get('custom_barcode') or '').strip()
    if qr.startswith('qr_images/') or qr.startswith('/qr_images/'):
        qr_url = f"https://app.bpmpjambi.my.id/siman/{qr.lstrip('/')}"
    elif qr.startswith('http://') or qr.startswith('https://') or qr.startswith('data:'):
        qr_url = qr
    elif qr:
        qr_url = f"https://app.bpmpjambi.my.id/siman/{qr.lstrip('/')}"
    else:
        qr_url = ""

    return f"""
    <article class="item-label">
      <header class="label-header">
        <div class="header-logo-box">
          <img src="{logo_base64}" alt="Logo" class="logo-image">
        </div>
        <div class="header-text-box">
          <div class="instansi-line1">KEMENTERIAN PENDIDIKAN DASAR DAN</div>
          <div class="instansi-line2">MENENGAH</div>
          <div class="satker-code-line">{html.escape(satker_display)}</div>
        </div>
        <div class="header-logo-box opacity-0 pointer-events-none" aria-hidden="true" style="visibility: hidden;"></div>
      </header>

      <div class="label-body">
        <div class="body-info-left">
          <div>
            <div class="info-top-row">
              <span class="info-kode-barang">{kode_barang}</span>
              <span class="info-nup">NUP: {nup}</span>
            </div>
            <div class="info-nama-barang">{nama_barang}</div>
          </div>
          <div class="info-merk-tipe">{merk_tipe}</div>
        </div>
        <div class="body-qr-right">
          <div class="qr-svg-wrapper">
            <img src="{qr_url}" alt="QR SIMAN" class="w-full h-full object-contain"
              onerror="this.style.display='none'; if(this.nextElementSibling) this.nextElementSibling.style.display='flex';"
              decoding="async">
            <div class="qr-fallback-msg" style="display:none; align-items:center; justify-content:center; width:100%; height:100%; font-size:7.5px; color:#94a3b8; font-family:'JetBrains Mono',monospace; text-align:center;">
              QR Belum Terdaftar
            </div>
          </div>
        </div>
      </div>
    </article>
    """

def get_page_head(title, subtitle=""):
    return f"""<!DOCTYPE html>
<html lang="id">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{html.escape(title)} - Cetak Label BMN A4</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600;700&display=swap" rel="stylesheet">
  <style>
    :root {{
      --page-width: 210mm;
      --page-height: 297mm;
      --page-margin-top: 13.71mm;
      --page-margin-bottom: 6.21mm;
      --page-margin-left: 0.67mm;
      --page-margin-right: 0.67mm;
      --label-width: 104mm;
      --label-height: 41mm;
      --header-height: 13.9mm;
      --label-col-gap: 0.67mm;
      --label-row-gap: 6.21mm;
      --grid-cols: 2;
      --qr-size: 21mm;
      --logo-size: 9mm;
      --instansi-font-size: 10.5pt;
      --satker-font-size: 9pt;
      --body-font-size: 9pt;
    }}
    * {{
      box-sizing: border-box;
      -webkit-print-color-adjust: exact !important;
      print-color-adjust: exact !important;
    }}
    body {{
      background-color: #0f172a;
      color: #0f172a;
      margin: 0;
      padding: 0;
      font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
      -webkit-font-smoothing: antialiased;
    }}
    .document-workspace {{
      display: flex;
      flex-direction: column;
      align-items: center;
      padding: 32px 0 64px;
      gap: 28px;
    }}
    .sheet-wrapper {{
      display: flex;
      flex-direction: column;
      align-items: center;
      margin-bottom: 24px;
    }}
    .a4-sheet {{
      width: var(--page-width);
      height: var(--page-height);
      min-width: var(--page-width);
      min-height: var(--page-height);
      max-width: var(--page-width);
      max-height: var(--page-height);
      background-color: #ffffff;
      padding-top: var(--page-margin-top);
      padding-bottom: var(--page-margin-bottom);
      padding-left: var(--page-margin-left);
      padding-right: var(--page-margin-right);
      box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.4);
      position: relative;
      overflow: hidden;
    }}
    .label-grid {{
      display: grid;
      grid-template-columns: repeat(2, var(--label-width));
      grid-auto-rows: var(--label-height);
      column-gap: var(--label-col-gap);
      row-gap: var(--label-row-gap);
      width: 100%;
      height: 100%;
      align-content: start;
      justify-content: center;
    }}
    .item-label {{
      width: var(--label-width);
      height: var(--label-height);
      border: 1px solid #000000;
      background: #ffffff;
      display: flex;
      flex-direction: column;
      overflow: hidden;
      position: relative;
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
    }}
    .label-header {{
      width: 100%;
      height: var(--header-height);
      min-height: var(--header-height);
      max-height: var(--header-height);
      border-bottom: 1px solid #000000;
      display: flex;
      flex-direction: row;
      align-items: center;
      padding: 0 2mm;
      gap: 1.5mm;
      overflow: hidden;
      flex-shrink: 0;
    }}
    .header-logo-box {{
      width: calc(var(--header-height) * 0.68);
      height: calc(var(--header-height) * 0.68);
      max-width: 14mm;
      max-height: 14mm;
      display: flex;
      align-items: center;
      justify-content: center;
      flex-shrink: 0;
    }}
    .logo-image {{
      width: var(--logo-size);
      height: var(--logo-size);
      max-width: 100%;
      max-height: 100%;
      object-fit: contain;
      display: block;
    }}
    .header-text-box {{
      flex: 1;
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      text-align: center;
      line-height: 1.15;
    }}
    .instansi-line1 {{
      font-size: var(--instansi-font-size);
      font-weight: 700;
      color: #000000;
      text-transform: uppercase;
      letter-spacing: -0.2px;
      white-space: nowrap;
    }}
    .instansi-line2 {{
      font-size: var(--instansi-font-size);
      font-weight: 700;
      color: #000000;
      text-transform: uppercase;
      letter-spacing: -0.2px;
      white-space: nowrap;
    }}
    .satker-code-line {{
      font-size: var(--satker-font-size);
      font-weight: 700;
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      color: #000000;
      letter-spacing: 0.1px;
      white-space: nowrap;
    }}
    .label-body {{
      flex: 1;
      display: flex;
      flex-direction: row;
      padding: 1.5mm 2mm;
      gap: 2mm;
      overflow: hidden;
      justify-content: space-between;
    }}
    .body-info-left {{
      flex: 1;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      overflow: hidden;
    }}
    .info-top-row {{
      display: flex;
      align-items: baseline;
      gap: 3mm;
      margin-bottom: 0.5mm;
    }}
    .info-kode-barang {{
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      font-size: var(--body-font-size);
      font-weight: 700;
      color: #000000;
      letter-spacing: 0.2px;
      line-height: 1.1;
    }}
    .info-nup {{
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      font-size: var(--body-font-size);
      font-weight: 700;
      color: #000000;
      line-height: 1.1;
    }}
    .info-nama-barang {{
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      font-size: var(--body-font-size);
      font-weight: 700;
      color: #000000;
      text-transform: uppercase;
      line-height: 1.15;
      display: -webkit-box;
      -webkit-line-clamp: 2;
      -webkit-box-orient: vertical;
      overflow: hidden;
    }}
    .info-merk-tipe {{
      font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
      font-size: 7.5pt;
      font-weight: 400;
      color: #334155;
      text-transform: uppercase;
      line-height: 1.1;
      white-space: nowrap;
      overflow: hidden;
      text-overflow: ellipsis;
    }}
    .body-qr-right {{
      width: var(--qr-size);
      height: var(--qr-size);
      display: flex;
      align-items: center;
      justify-content: center;
      align-self: center;
      flex-shrink: 0;
    }}
    .qr-svg-wrapper {{
      width: 100%;
      height: 100%;
      display: flex;
      align-items: center;
      justify-content: center;
    }}
    .sheet-paper-header {{
      position: absolute;
      top: 4.0mm;
      left: var(--page-margin-left);
      right: var(--page-margin-right);
      display: flex;
      align-items: center;
      justify-content: center;
      font-family: 'JetBrains Mono', monospace;
      font-size: 8pt;
      font-weight: 700;
      color: #0284c7;
      text-transform: uppercase;
      letter-spacing: 0.5px;
      pointer-events: none;
    }}
    .sheet-paper-single-title {{
      display: inline-flex;
      align-items: center;
      gap: 6px;
      background: #f8fafc;
      padding: 2px 12px;
      border-radius: 4px;
      border: 1px solid #e2e8f0;
    }}
    .sheet-room-tag {{ color: #0369a1; font-weight: 800; }}
    .sheet-room-name {{ color: #0f172a; font-weight: 800; }}
    .sheet-header-bullet {{ color: #94a3b8; font-weight: 400; }}
    .sheet-page-info {{ color: #475569; }}
    .sheet-count-info {{ color: #64748b; font-size: 7.5pt; font-weight: 500; }}
    @media print {{
      @page {{
        size: A4 portrait;
        margin: 0;
      }}
      body {{ background-color: transparent !important; }}
      .sheet-wrapper {{
        margin: 0 !important;
        page-break-after: always !important;
        break-after: page !important;
      }}
      .a4-sheet {{
        box-shadow: none !important;
        page-break-inside: avoid !important;
        break-inside: avoid !important;
      }}
    }}
  </style>
</head>
<body>
  <div class="document-workspace">
"""

def get_page_foot():
    return """
  </div>
</body>
</html>
"""

def generate_room_html(room_name, items):
    LABELS_PER_PAGE = 12
    total_sheets = (len(items) + LABELS_PER_PAGE - 1) // LABELS_PER_PAGE
    room_pages_html = []

    for sheet_idx in range(total_sheets):
        start = sheet_idx * LABELS_PER_PAGE
        sheet_items = items[start:start + LABELS_PER_PAGE]
        labels_html = "\n".join([render_label_html(it) for it in sheet_items])
        sheet_block = f"""
        <div class="sheet-wrapper">
          <div class="a4-sheet">
            <div class="sheet-paper-header">
              <div class="sheet-paper-single-title">
                <span class="sheet-room-tag">RUANGAN :</span>
                <span class="sheet-room-name">{html.escape(room_name)}</span>
                <span class="sheet-header-bullet">&bull;</span>
                <span class="sheet-page-info">Halaman {sheet_idx + 1} dari {total_sheets}</span>
                <span class="sheet-count-info">({len(sheet_items)} Label BMN)</span>
              </div>
            </div>
            <div class="label-grid">
              {labels_html}
            </div>
          </div>
        </div>
        """
        room_pages_html.append(sheet_block)

    return get_page_head(room_name, f"{len(items)} Label") + "\n".join(room_pages_html) + get_page_foot()

def main():
    print("=" * 70)
    print(">>> GENERATOR PERBAIKAN TAHUN LABEL PDF RUANGAN")
    print("=" * 70)

    excel_path = os.path.join(DATA_DIR, "data.xlsx")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb['Database'] if 'Database' in wb.sheetnames else wb.active

    # Daftar target ruangan yang perlu diperbarui
    # Jika parameter target diberikan di CLI: python scripts/fix_rooms_pdf.py [target...]
    # Default: update [I.1.04] Ruang Rapat Angso Duo dan [J.1.04] Gudang
    target_args = sys.argv[1:] if len(sys.argv) > 1 else ['angso', 'j.1.04']
    print(f"Target filter: {target_args}")

    rooms_data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 6: continue
        kd_r = str(row[20] or '').strip()
        nm_r = str(row[21] or '').strip()
        bg = get_building_name(kd_r, nm_r)

        if bg:
            room_title = f"[{kd_r}] {nm_r} - {bg}" if kd_r and nm_r else (f"[{kd_r}] {bg}" if kd_r else f"{nm_r} - {bg}")
        elif kd_r and nm_r:
            room_title = f"[{kd_r}] {nm_r}"
        elif nm_r:
            room_title = nm_r
        elif kd_r:
            room_title = f"[{kd_r}] Ruangan {kd_r}"
        else:
            room_title = "Tanpa Ruangan"

        # Cek apakah cocok dengan salah satu target_args
        full_match_str = (kd_r + ' ' + nm_r + ' ' + room_title).lower()
        if not any(t.lower() in full_match_str for t in target_args):
            continue

        satker = str(row[1] or '138041000693242000KD').strip()
        tgl_perolehan = str(row[9] or '').strip() if len(row) > 9 and row[9] is not None else ''
        tahun = extract_year(row[25] if len(row) > 25 else '', row[9] if len(row) > 9 else '', satker)

        item = {
            'kode_satker': satker,
            'tahun': tahun,
            'tgl_perolehan': tgl_perolehan,
            'kode_barang': str(row[3] or '-').strip(),
            'nup': str(row[4] or '1').strip(),
            'nama_barang': str(row[5] or 'Barang BMN').strip(),
            'merk_tipe': str(row[8] or '-').strip() if len(row) > 8 and row[8] else '-',
            'custom_barcode': str(row[27] or '').strip() if len(row) > 27 else '',
            'ruangan': room_title
        }

        folder_name = sanitize_filename(room_title, 65)
        if folder_name not in rooms_data:
            rooms_data[folder_name] = {'title': room_title, 'items': []}
        rooms_data[folder_name]['items'].append(item)

    wb.close()
    print(f"[*] Ditemukan {len(rooms_data)} ruangan target dengan total {sum(len(v['items']) for v in rooms_data.values())} label:")
    for f_name, r_info in rooms_data.items():
        print(f"   -> {f_name}: {len(r_info['items'])} items")

        # Buat folder dan simpan HTML
        room_folder_path = os.path.join(OUTPUT_DIR, f_name)
        os.makedirs(room_folder_path, exist_ok=True)

        html_content = generate_room_html(r_info['title'], r_info['items'])
        html_file = os.path.join(room_folder_path, "template_print.html")
        with open(html_file, "w", encoding="utf-8") as f:
            f.write(html_content)
        print(f"      [OK] Tersimpan HTML: {html_file}")

    print("\n[OK] Penulisan HTML selesai. Sekarang menjalankan render Chrome CDP via scripts/03_render_pdfs.js...")
    subprocess.run(["node", os.path.join(BASE_DIR, "scripts", "03_render_pdfs.js")])

if __name__ == "__main__":
    main()
