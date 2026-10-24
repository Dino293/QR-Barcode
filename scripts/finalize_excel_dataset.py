import os
import re
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
EXCEL_PATH = os.path.join(DATA_DIR, 'data.xlsx')

def evaluate_year(t_val, tgl_val, satker_val):
    t_str = str(t_val or '').strip()
    if not t_str.startswith('='):
        m = re.search(r'\b(19\d{2}|20\d{2})\b', t_str)
        if m:
            return int(m.group(1))
            
    d_str = str(tgl_val or '').strip()
    m2 = re.search(r'\b(19\d{2}|20\d{2})\b', d_str)
    if m2:
        return int(m2.group(1))
        
    s_str = str(satker_val or '').strip()
    m3 = re.search(r'\.(19\d{2}|20\d{2})\b', s_str)
    if m3:
        return int(m3.group(1))
        
    return 2024

def finalize():
    print("Membaca workbook untuk standarisasi Kolom Tahun...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Database'] if 'Database' in wb.sheetnames else wb.active
    
    updated_years = 0
    for r in range(2, ws.max_row + 1):
        kd = ws.cell(row=r, column=4).value
        if not kd:
            continue
            
        cur_t = ws.cell(row=r, column=26).value
        tgl = ws.cell(row=r, column=10).value
        satker = ws.cell(row=r, column=2).value
        
        yr = evaluate_year(cur_t, tgl, satker)
        if str(cur_t).startswith('=') or cur_t is None or cur_t != yr:
            ws.cell(row=r, column=26, value=yr)
            updated_years += 1
            
    print(f"Total baris tahun distandarisasi ke nilai integer: {updated_years:,}")
    print("Menyimpan perubahan...")
    wb.save(EXCEL_PATH)
    wb.close()
    print("Selesai standarisasi kolom Tahun!")

if __name__ == '__main__':
    finalize()
