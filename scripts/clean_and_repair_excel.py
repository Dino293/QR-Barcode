import os
import sys
import shutil
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
EXCEL_PATH = os.path.join(DATA_DIR, 'data.xlsx')
CACHE_PATH = os.path.join(DATA_DIR, 'siman_cache.json')
MAP_PATH = os.path.join(DATA_DIR, 'clean_room_building_map.json')
BACKUP_DIR = os.path.join(DATA_DIR, 'backup')

def clean_and_repair():
    print("=" * 80)
    print("MEMBERSIHKAN DAN MEMPERBAIKI DATASET EXCEL (data.xlsx)")
    print("=" * 80)
    
    # 1. Pastikan folder backup ada dan simpan backup
    os.makedirs(BACKUP_DIR, exist_ok=True)
    backup_path = os.path.join(BACKUP_DIR, 'data_backup_before_clean.xlsx')
    shutil.copy2(EXCEL_PATH, backup_path)
    print(f"[*] Backup tersimpan di: {backup_path}")
    
    # 2. Muat SIMAN Cache
    with open(CACHE_PATH, 'r', encoding='utf-8') as f:
        cache = json.load(f)
    print(f"[*] SIMAN Cache: {len(cache):,} data")
    
    # 3. Muat Building Map
    with open(MAP_PATH, 'r', encoding='utf-8') as f:
        mapping = json.load(f)
    pairs = mapping.get('pairs', {})
    codes = mapping.get('codes', {})
    
    def get_building_name(kd, nm):
        kd_u = str(kd or '').strip().upper()
        nm_l = str(nm or '').strip().lower()
        if not kd_u and not nm_l:
            return ''
        key = f"{kd_u.lower()}|||{nm_l}"
        if key in pairs:
            return pairs[key]
        if kd_u.startswith('D.1.') or kd_u.startswith('D.2.') or kd_u.startswith('D.'):
            return 'Asrama Marsawa'
        if kd_u.startswith('F.1.') or kd_u.startswith('F.2.') or kd_u.startswith('F.'):
            return 'Asrama Bulian'
        if kd_u.startswith('E.1.') or kd_u.startswith('E.'):
            return 'Asrama Kulim'
        if kd_u.startswith('G.1.') or kd_u.startswith('G.2.') or kd_u.startswith('G.'):
            return 'Mess Jelutung 3'
        if kd_u.startswith('H.1.') or kd_u.startswith('H.'):
            return 'Mess Penatar'
        if (kd_u == '2' or kd_u == '3') and 'kamar' in nm_l:
            return 'Mess Kepala'
        if (kd_u == '4' or kd_u == '5') and 'kamar' in nm_l:
            return 'Mess Jelutung 2'
        kd_l = kd_u.lower()
        if kd_l not in ['1', '2', '3', '4', '5', '01', '02', '03', '04', '05'] and kd_l in codes:
            return codes[kd_l]
        return ''

    # 4. Buka Workbook dalam mode tulis
    print(f"[*] Membuka workbook '{EXCEL_PATH}'...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Database'] if 'Database' in wb.sheetnames else wb.active
    
    # Pastikan Header Baris 1:
    # Kolom 9 (I): Merk/Tipe
    # Kolom 28 (AB): QR
    # Kolom 29 (AC): Nama Bangunan
    if ws.cell(1, 9).value is None:
        ws.cell(1, 9, "Merk/Tipe")
    if ws.cell(1, 28).value is None:
        ws.cell(1, 28, "QR")
    if ws.cell(1, 29).value is None:
        ws.cell(1, 29, "Nama Bangunan")
        
    total_processed = 0
    merk_repaired = 0
    building_repaired = 0
    qr_synced = 0
    satker_fixed = 0
    
    for row_idx in range(2, ws.max_row + 1):
        kd_brg = ws.cell(row=row_idx, column=4).value
        if kd_brg is None or str(kd_brg).strip() == '':
            continue
            
        total_processed += 1
        kd_brg = str(kd_brg).strip()
        nup = str(ws.cell(row=row_idx, column=5).value or '1').strip()
        
        # 1. Satker
        satker_val = ws.cell(row=row_idx, column=2).value
        if not satker_val or str(satker_val).strip() == '':
            ws.cell(row=row_idx, column=2, value="138041000693242000KD")
            satker_fixed += 1
            
        # 2. Merk / Tipe
        merk_val = ws.cell(row=row_idx, column=9).value
        if merk_val is None or str(merk_val).strip() in ('', 'None', 'null'):
            ws.cell(row=row_idx, column=9, value="-")
            merk_repaired += 1
            
        # 3. QR dari SIMAN Cache jika kosong
        qr_val = ws.cell(row=row_idx, column=28).value
        cache_key = f"{kd_brg}_{nup}"
        siman_qr = cache.get(cache_key, "")
        if (not qr_val or str(qr_val).strip() == '') and siman_qr:
            ws.cell(row=row_idx, column=28, value=siman_qr)
            qr_synced += 1
            
        # 4. Nama Bangunan (Kolom 29 / AC)
        kd_r = str(ws.cell(row=row_idx, column=21).value or '').strip()
        nm_r = str(ws.cell(row=row_idx, column=22).value or '').strip()
        curr_b = str(ws.cell(row=row_idx, column=29).value or '').strip()
        canonical_b = get_building_name(kd_r, nm_r)
        
        if curr_b != canonical_b:
            ws.cell(row=row_idx, column=29, value=canonical_b)
            building_repaired += 1
            
    print(f"[*] Menyimpan perubahan ke '{EXCEL_PATH}'...")
    wb.save(EXCEL_PATH)
    wb.close()
    
    print("-" * 80)
    print("HASIL PERBAIKAN DATASET:")
    print("-" * 80)
    print(f"Total baris diproses         : {total_processed:,}")
    print(f"Merk/Tipe kosong diisi '-'   : {merk_repaired:,}")
    print(f"Nama Bangunan disinkronisasi : {building_repaired:,}")
    print(f"QR Barcode baru diisi        : {qr_synced:,}")
    print(f"Satker kosong diperbaiki     : {satker_fixed:,}")
    print("=" * 80)
    print("[SUKSES] Dataset Excel berhasil direvisi & distandarisasi 100%!")

if __name__ == '__main__':
    clean_and_repair()
