import os
import sys
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
EXCEL_PATH = os.path.join(DATA_DIR, 'data.xlsx')
CACHE_PATH = os.path.join(DATA_DIR, 'siman_cache.json')
MAP_PATH = os.path.join(DATA_DIR, 'clean_room_building_map.json')

def run_crosscheck():
    print("=" * 80)
    print("EVALUASI DAN CROSSCHECK DATA EXCEL VS SIMAN CACHE")
    print("=" * 80)
    
    # Load SIMAN cache
    with open(CACHE_PATH, 'r', encoding='utf-8') as f:
        cache = json.load(f)
    print(f"Total entries in SIMAN cache: {len(cache):,}")
    
    # Load Room-Building Map
    with open(MAP_PATH, 'r', encoding='utf-8') as f:
        mapping = json.load(f)
    pairs = mapping.get('pairs', {})
    codes = mapping.get('codes', {})
    
    def get_building_name(kd, nm):
        kd_u = str(kd or '').strip().upper()
        nm_l = str(nm or '').strip().lower()
        if not kd_u and not nm_l:
            return ''
        key = f"{kd_u.lower()}|||{nm_l}"
        if key in pairs:
            return pairs[key]
        if kd_u.startswith('D.1.') or kd_u.startswith('D.2.') or kd_u.startswith('D.'):
            return 'Asrama Marsawa'
        if kd_u.startswith('F.1.') or kd_u.startswith('F.2.') or kd_u.startswith('F.'):
            return 'Asrama Bulian'
        if kd_u.startswith('E.1.') or kd_u.startswith('E.'):
            return 'Asrama Kulim'
        if kd_u.startswith('G.1.') or kd_u.startswith('G.2.') or kd_u.startswith('G.'):
            return 'Mess Jelutung 3'
        if kd_u.startswith('H.1.') or kd_u.startswith('H.'):
            return 'Mess Penatar'
        if (kd_u == '2' or kd_u == '3') and 'kamar' in nm_l:
            return 'Mess Kepala'
        if (kd_u == '4' or kd_u == '5') and 'kamar' in nm_l:
            return 'Mess Jelutung 2'
        kd_l = kd_u.lower()
        if kd_l not in ['1', '2', '3', '4', '5', '01', '02', '03', '04', '05'] and kd_l in codes:
            return codes[kd_l]
        return ''

    # Load Excel
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb['Database'] if 'Database' in wb.sheetnames else wb.active
    
    headers = []
    rows_data = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            headers = [str(c or '') for c in row]
            continue
        if not any(row):
            continue
        rows_data.append(row)
    wb.close()
    
    total_rows = len(rows_data)
    print(f"Total baris data di Excel: {total_rows:,}")
    print(f"Headers ({len(headers)}): {headers}\n")
    
    # Column indices:
    # 1: Kode Satker, 2: Uraian Satker, 3: Kode Barang, 4: NUP, 5: Nama Barang, 8: Merk/Tipe, 9: Tgl Perolehan,
    # 20: Kode Ruangan, 21: Nama Ruangan, 25: Tahun, 27: QR, 28: Nama Bangunan
    
    stats = {
        'empty_satker': [],
        'empty_tahun': [],
        'empty_nup': [],
        'empty_nama_barang': [],
        'empty_kode_barang': [],
        'empty_merk_tipe': [],
        'empty_qr': [],
        'empty_kode_ruangan': [],
        'empty_nama_ruangan': [],
        'empty_nama_bangunan_col': [],
        'unmapped_building_func': [],
        'qr_mismatch_with_siman': [],
        'missing_in_siman_cache': [],
        'invalid_nup': [],
        'satker_values': set(),
        'year_values': set()
    }
    
    for r_idx, row in enumerate(rows_data, start=2):
        satker = str(row[1] or '').strip()
        kd_brg = str(row[3] or '').strip()
        nup = str(row[4] or '').strip()
        nm_brg = str(row[5] or '').strip()
        merk = str(row[8] or '').strip() if len(row) > 8 else ''
        tgl = str(row[9] or '').strip() if len(row) > 9 else ''
        kd_r = str(row[20] or '').strip() if len(row) > 20 else ''
        nm_r = str(row[21] or '').strip() if len(row) > 21 else ''
        tahun = str(row[25] or '').strip() if len(row) > 25 else ''
        qr = str(row[27] or '').strip() if len(row) > 27 else ''
        nm_b = str(row[28] or '').strip() if len(row) > 28 else ''
        
        # Check Satker
        if not satker:
            stats['empty_satker'].append(r_idx)
        else:
            stats['satker_values'].add(satker)
            
        # Check Tahun
        if not tahun:
            stats['empty_tahun'].append((r_idx, tgl))
        else:
            stats['year_values'].add(tahun)
            
        # Check NUP
        if not nup:
            stats['empty_nup'].append(r_idx)
        else:
            try:
                int(float(nup))
            except:
                stats['invalid_nup'].append((r_idx, nup))
                
        # Check Nama Barang
        if not nm_brg:
            stats['empty_nama_barang'].append(r_idx)
            
        # Check Kode Barang
        if not kd_brg:
            stats['empty_kode_barang'].append(r_idx)
            
        # Check Merk/Tipe
        if not merk or merk == '-' or merk.lower() == 'none':
            stats['empty_merk_tipe'].append((r_idx, nm_brg))
            
        # Check QR & SIMAN
        cache_key = f"{kd_brg}_{nup}"
        siman_qr = cache.get(cache_key)
        
        if not qr:
            stats['empty_qr'].append((r_idx, kd_brg, nup))
        else:
            if siman_qr is None:
                stats['missing_in_siman_cache'].append((r_idx, kd_brg, nup, qr))
            elif qr != siman_qr:
                stats['qr_mismatch_with_siman'].append((r_idx, kd_brg, nup, qr, siman_qr))
                
        # Check Kode Ruangan & Nama Ruangan
        if not kd_r:
            stats['empty_kode_ruangan'].append(r_idx)
        if not nm_r:
            stats['empty_nama_ruangan'].append(r_idx)
            
        # Check Nama Bangunan di Kolom 28 (AC)
        if not nm_b:
            stats['empty_nama_bangunan_col'].append(r_idx)
            
        # Check mapping function building name
        bg_func = get_building_name(kd_r, nm_r)
        if not bg_func:
            stats['unmapped_building_func'].append((kd_r, nm_r))

    print("-" * 80)
    print("HASIL CROSSCHECK RINGKAS:")
    print("-" * 80)
    print(f"1. Kode Satker:")
    print(f"   - Kosong: {len(stats['empty_satker'])}")
    print(f"   - Distinct Values: {stats['satker_values']}")
    
    print(f"\n2. Tahun:")
    print(f"   - Kosong: {len(stats['empty_tahun'])}")
    print(f"   - Distinct Year Count: {len(stats['year_values'])}")
    sample_years = sorted(list(stats['year_values']))
    print(f"   - Rentang Tahun: {sample_years[0]} s/d {sample_years[-1]}")
    
    print(f"\n3. NUP:")
    print(f"   - Kosong: {len(stats['empty_nup'])}")
    print(f"   - Invalid format: {len(stats['invalid_nup'])}")
    
    print(f"\n4. Nama Barang:")
    print(f"   - Kosong: {len(stats['empty_nama_barang'])}")
    
    print(f"\n5. Kode Barang:")
    print(f"   - Kosong: {len(stats['empty_kode_barang'])}")
    
    print(f"\n6. Merek/Tipe Barang:")
    print(f"   - Kosong/Strip: {len(stats['empty_merk_tipe'])} ({(len(stats['empty_merk_tipe'])/total_rows)*100:.2f}%)")
    
    print(f"\n7. QR Barcode & SIMAN Match:")
    print(f"   - Kosong di Excel: {len(stats['empty_qr'])}")
    print(f"   - Tidak ada di SIMAN Cache: {len(stats['missing_in_siman_cache'])}")
    print(f"   - Beda nilai Excel vs SIMAN: {len(stats['qr_mismatch_with_siman'])}")
    
    print(f"\n8. Kode Ruangan:")
    print(f"   - Kosong: {len(stats['empty_kode_ruangan'])}")
    
    print(f"\n9. Nama Ruangan:")
    print(f"   - Kosong: {len(stats['empty_nama_ruangan'])}")
    
    print(f"\n10. Nama Bangunan / Ruangan Gabungan:")
    print(f"   - Kolom AC (Nama Bangunan di Excel) Kosong: {len(stats['empty_nama_bangunan_col'])} ({(len(stats['empty_nama_bangunan_col'])/total_rows)*100:.2f}%)")
    
    distinct_unmapped = set(stats['unmapped_building_func'])
    print(f"   - Ruangan yang belum terpetakan ke Gedung (unmapped): {len(distinct_unmapped)} ruangan unik")
    if distinct_unmapped:
        print("     Daftar ruangan tanpa mapping gedung:")
        for kd, nm in sorted(distinct_unmapped):
            print(f"     * [{kd}] {nm}")
            
    print("=" * 80)

if __name__ == '__main__':
    run_crosscheck()
