"""
========================================================================================
SKRIP MIGRASI DATA EXCEL (data.xlsx) KE SUPABASE POSTGRESQL
Sistem Inventarisasi & Pencetakan Barcode BMN - BPMP Provinsi Jambi
========================================================================================
Tugas:
1. Membaca 14.359 data baris dari data/data.xlsx (Sheet 'Database').
2. Melakukan pembersihan data (cleaning & formatting).
3. Mengunggah secara efisien menggunakan batch upsert (500 data per batch) ke tabel 'bmn_items'.
========================================================================================
"""

import sys
import os
import time
import openpyxl
from supabase import create_client, Client

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
env_path = os.path.join(BASE_DIR, '.env')
if os.path.exists(env_path):
    try:
        with open(env_path, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    os.environ.setdefault(k.strip(), v.strip())
    except Exception:
        pass

# 1. Konfigurasi Kredensial Supabase
SUPABASE_URL = os.getenv("SUPABASE_URL", "https://vipdzgwzxvctofcwityn.supabase.co")
SUPABASE_KEY = os.getenv("SUPABASE_KEY", "")
DATA_FILE = os.path.join(BASE_DIR, "data", "data.xlsx")
BATCH_SIZE = 500


def bersihkan_angka(val, default=0):
    if val is None or val == "":
        return default
    try:
        return int(float(str(val).replace(",", ".").strip()))
    except Exception:
        return default


def migrate_data(supabase_key=None):
    key = supabase_key or SUPABASE_KEY
    if not key:
        print("[ERROR] API Key Supabase belum diisi!")
        print("Silakan jalankan: python scripts/migrate_to_supabase.py <SUPABASE_KEY>")
        return False

    print("=" * 70)
    print(">>> MEMULAI MIGRASI DATA EXCEL KE CLOUD SUPABASE (PostgreSQL)")
    print(f"[*] Target Supabase : {SUPABASE_URL}")
    print(f"[*] Berkas Sumber   : {DATA_FILE}")
    print(f"[*] Ukuran Batch    : {BATCH_SIZE} baris per batch")
    print("=" * 70)

    try:
        supabase: Client = create_client(SUPABASE_URL, key)
    except Exception as e:
        print(f"[ERROR] Gagal inisialisasi koneksi Supabase: {str(e)}")
        return False

    if not os.path.exists(DATA_FILE):
        print(f"[ERROR] Berkas tidak ditemukan: {DATA_FILE}")
        return False

    print("\n[1/3] Membuka dan membaca data.xlsx...")
    wb = openpyxl.load_workbook(DATA_FILE, read_only=True, data_only=True)
    if "Database" in wb.sheetnames:
        ws = wb["Database"]
    else:
        ws = wb.active

    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)

    if not headers:
        print("[ERROR] Header sheet kosong!")
        wb.close()
        return False

    records = []
    skipped = 0

    print("[2/3] Memproses dan memvalidasi baris data...")
    for idx, row in enumerate(rows, start=1):
        if not row or len(row) < 5:
            skipped += 1
            continue

        kode_barang = str(row[3] or "").strip()
        if not kode_barang or kode_barang == "None":
            skipped += 1
            continue

        nup = bersihkan_angka(row[4], default=1)
        nama_barang = str(row[5] or "Barang BMN").strip()
        tgl_perolehan = str(row[9] or "").strip()
        tahun = bersihkan_angka(row[25], default=None)

        if not tahun and tgl_perolehan and len(tgl_perolehan) >= 4:
            try:
                tahun = int(tgl_perolehan[:4])
            except Exception:
                pass

        qr_val = str(row[27] or "").strip() if len(row) > 27 else ""
        nm_bangunan = str(row[28] or "").strip() if len(row) > 28 else ""

        record = {
            "no_urut": bersihkan_angka(row[0], default=idx),
            "kode_satker": str(row[1] or "138041000693242000KD").strip(),
            "uraian_satker": str(row[2] or "BPMP PROVINSI JAMBI").strip(),
            "kode_barang": kode_barang,
            "nup": nup,
            "nama_barang": nama_barang,
            "merk_tipe": str(row[8] or "-").strip(),
            "tgl_perolehan": tgl_perolehan,
            "tahun": tahun,
            "kuantitas": bersihkan_angka(row[10], default=1),
            "kondisi": str(row[11] or "").strip(),
            "harga_barang": str(row[12] or "").strip(),
            "status_penggunaan": str(row[13] or "").strip(),
            "kode_ruangan": str(row[20] or "").strip(),
            "nama_ruangan": str(row[21] or "").strip(),
            "jenis_bangunan": str(row[22] or "").strip(),
            "status_kodefikasi": str(row[23] or "").strip(),
            "status_inventarisasi": str(row[24] or "").strip(),
            "keterangan": str(row[26] or "").strip() if len(row) > 26 else "",
            "custom_barcode": qr_val,
            "nama_bangunan": nm_bangunan,
        }
        records.append(record)

    wb.close()
    total_valid = len(records)
    print(f"[*] Berhasil memvalidasi {total_valid:,} baris data (dilewati: {skipped})")

    print(f"\n[3/3] Mengunggah ke tabel 'bmn_items' di Supabase...")
    start_time = time.time()
    uploaded_count = 0

    for i in range(0, total_valid, BATCH_SIZE):
        batch = records[i:i + BATCH_SIZE]
        batch_no = (i // BATCH_SIZE) + 1
        total_batches = (total_valid + BATCH_SIZE - 1) // BATCH_SIZE

        try:
            # Gunakan upsert dengan on_conflict kode_barang,nup
            response = supabase.table("bmn_items").upsert(
                batch,
                on_conflict="kode_barang,nup"
            ).execute()

            uploaded_count += len(batch)
            pct = (uploaded_count / total_valid) * 100
            print(f"   -> Batch {batch_no}/{total_batches}: +{len(batch)} data terunggah ({pct:.1f}%)")

        except Exception as e:
            print(f"   [!] Gagal pada batch {batch_no}: {str(e)}")
            time.sleep(1)

    duration = round(time.time() - start_time, 2)
    print("\n" + "=" * 70)
    print(">>> MIGRASI KE SUPABASE SELESAI!")
    print(f"[*] Total Baris Terunggah : {uploaded_count:,} / {total_valid:,}")
    print(f"[*] Waktu Eksekusi        : {duration} detik")
    print("=" * 70)
    return True


if __name__ == "__main__":
    key_arg = sys.argv[1] if len(sys.argv) > 1 else None
    migrate_data(supabase_key=key_arg)
